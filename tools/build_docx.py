#!/usr/bin/env python3
"""build_docx.py — formatted analytical report (.docx) from a JSON (or Markdown) content spec.

Features: title page, auto table of contents (Word field with pre-filled entries), headings,
paragraphs with **bold**/*italic*, bullet & numbered lists (numbering restarts per list),
tables with captions and source lines, NATIVE EDITABLE WORD CHARTS (DrawingML chart part +
embedded .xlsx workbook, "Edit Data" works in Word), key-figure panels, call-outs, scorecards,
real Word footnotes (GOST-style source citations) or numeric [n] references, bibliography,
header/footer with page numbers. Font is configurable (default Times New Roman).

JSON spec:
{
  "meta": {"title": "...", "subtitle": "...", "org": "...", "author": "...", "date": "...",
           "language": "ru", "font": "Times New Roman", "font_size": 12, "heading_font": null,
           "preset": "report" | "gost",          # gost = 14 pt, 1.5 spacing, 1.25 cm indent, GOST 7.32 margins
           "citation_style": "footnotes" | "numeric" | "author_date", "accent": "0F766E",
           "watermark": "ПРИМЕР · вымышленные данные", "facts": "facts.json", "theme": "light"},
  "sources": [{"id": "s1", "gost": "Full GOST R 7.0.100-2018 style reference", "url": "...", "accessed": "20.10.2026"}],
  "blocks": [
    {"type": "title_page"}, {"type": "toc"}, {"type": "page_break"},
    {"type": "heading", "level": 1, "text": "1. Резюме"},
    {"type": "paragraph", "text": "Рынок вырастет до {{market.tam_2030|money_bn}} [[src:s1]]."},
    {"type": "bullets", "items": ["...", "..."]}, {"type": "numbered", "items": ["..."]},
    {"type": "key_figures", "items": [{"value": "12,4 млрд ₽", "label": "SAM 2030"}], "source": "..."},
    {"type": "callout", "title": "Вывод", "text": "..."},
    {"type": "table", "caption": "...", "columns": [...], "rows": [[...]], "source": "...", "col_widths_cm": [...]},
    {"type": "chart", "caption": "...", "chart": {<charts_common spec>}, "source": "...", "width_cm": 16, "height_cm": 8},
    {"type": "scorecard", "caption": "...", "criteria": [{"name": "...", "weight": 14, "score": 3.5, "note": ""}], "total": 66},
    {"type": "sources"},           # bibliography
    {"type": "glossary", "ids": ["tam", "sam"]}   # "What the numbers mean"; inserted automatically before
                                   # the first heading with the metrics the report shows (config/glossary.json);
                                   # meta.glossary = false switches it off
  ]
}
Inline citations: [[src:<id>]] → footnote (or [n]); [[fn:free text]] → free footnote.
Markdown input (--markdown file.md): '#'-headings, paragraphs, '-' bullets, '1.' lists, pipe tables,
```chart {json}``` blocks, '[TOC]', '\\pagebreak'.

CLI: python build_docx.py --input report_spec.json --out report.docx [--facts facts.json] [--font "PT Serif"]
     python build_docx.py --markdown report.md --meta meta.json --out report.docx
"""
from __future__ import annotations

import argparse
import copy
import json
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Emu, Pt, RGBColor
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import glossary  # noqa: E402
from charts_common import add_ea_fonts, fix_axis_ids, prepare, set_chart_lang, style_chart  # noqa: E402
from ooxml_clean import clean_package  # noqa: E402
from fmt_utils import dump_json, fmt_number, load_json, resolve_placeholders  # noqa: E402
from citations import format_reference, inline_citation, sort_key, STYLE_FOR_LOCALE, INLINE_FOR_STYLE  # noqa: E402
from locales import disclaimer, disclaimer_title, fmt_date, get_locale, locale_of  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
THEMES_FILE = os.path.join(HERE, "..", "config", "themes.json")

def _labels(code):
    return get_locale(code)["labels"]


CT_CHART = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
CT_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CT_FOOTNOTES = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
RT_PACKAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _rgb(h):
    return RGBColor.from_string(h.replace("#", "").upper())


SETTINGS_TAIL = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                 "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                 "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                 "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]
TBLPR_AFTER_BORDERS = ["shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]


def insert_before(parent, el, successor_tags):
    """Insert el before the first existing child whose local name is in successor_tags (schema order)."""
    for tag in successor_tags:
        found = parent.find(qn(f"w:{tag}"))
        if found is not None:
            found.addprevious(el)
            return el
    parent.append(el)
    return el


def set_num(paragraph, num_id, ilvl=0):
    numPr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = ilvl
    numPr.get_or_add_numId().val = num_id


# ----------------------------------------------------------------------------- markdown → blocks
def markdown_to_blocks(md: str) -> list[dict]:
    blocks, lines, i = [], md.splitlines(), 0
    para: list[str] = []

    def flush():
        if para:
            blocks.append({"type": "paragraph", "text": " ".join(x.strip() for x in para)})
            para.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            flush()
            kind = s[3:].strip() or "code"
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            body = "\n".join(buf)
            if kind in ("chart", "block", "json"):
                obj = json.loads(body)
                blocks.append(obj if kind != "chart" else {"type": "chart", **obj} if "chart" in obj
                              else {"type": "chart", "chart": obj})
            else:
                blocks.append({"type": "paragraph", "text": body})
        elif not s:
            flush()
        elif s in ("[TOC]", "[[TOC]]"):
            flush(); blocks.append({"type": "toc"})
        elif s in ("\\pagebreak", "\\newpage", "---pagebreak---"):
            flush(); blocks.append({"type": "page_break"})
        elif re.match(r"^#{1,4}\s", s):
            flush()
            lvl = len(s) - len(s.lstrip("#"))
            blocks.append({"type": "heading", "level": lvl, "text": s[lvl:].strip()})
        elif re.match(r"^[-*]\s", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s", lines[i]):
                lv = 1 if lines[i].startswith(("  ", "\t")) else 0
                items.append({"text": re.sub(r"^\s*[-*]\s", "", lines[i]), "level": lv})
                i += 1
            blocks.append({"type": "bullets", "items": items})
            continue
        elif re.match(r"^\d+[.)]\s", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*\d+[.)]\s", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s", "", lines[i]))
                i += 1
            blocks.append({"type": "numbered", "items": items})
            continue
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            cap = None
            if blocks and blocks[-1]["type"] == "paragraph" and re.match(r"^(Таблица|Table|表)\s*[:.：]", blocks[-1]["text"]):
                cap = re.sub(r"^(Таблица|Table|表)\s*[:.：]\s*", "", blocks.pop()["text"])
            blocks.append({"type": "table", "columns": rows[0], "rows": rows[1:], "caption": cap})
            continue
        else:
            para.append(s)
        i += 1
    flush()
    return blocks


# ----------------------------------------------------------------------------- builder
class ReportBuilder:
    def __init__(self, spec: dict):
        self.spec = spec
        self.meta = spec.get("meta", {})
        self.lang = locale_of(self.meta)
        self.loc = get_locale(self.lang)
        self.L = _labels(self.lang)
        themes = load_json(THEMES_FILE)
        self.theme = copy.deepcopy(themes.get(self.meta.get("theme", "light"), themes["light"]))
        self.accent = self.meta.get("accent", self.theme["accent"])
        F = self.loc["fonts"]
        self.font = self.meta.get("font") or F["report_latin"]
        self.hfont = self.meta.get("heading_font") or (F["report_heading"] if not self.meta.get("font") else self.font)
        self.ea_font = self.meta.get("font_ea") or F.get("report_ea") or self.font
        self.hea_font = self.meta.get("heading_font_ea") or F.get("report_heading_ea") or self.ea_font
        self.ea_map = {self.hfont: self.hea_font} if self.hfont != self.font else {}
        preset = self.meta.get("preset", "report")
        self.gost = preset == "gost"
        self.size = self.meta.get("font_size", 14 if self.gost else 12)
        self.spacing = self.meta.get("line_spacing", 1.5 if self.gost else 1.15)
        chars = self.loc.get("first_line_indent_chars", 0)
        self.indent_cm = self.meta.get("first_line_indent_cm", 1.25 if self.gost else chars * self.size * 2.54 / 72)
        self.bib_style = self.meta.get("bibliography_style") or STYLE_FOR_LOCALE.get(self.lang, "apa")
        self.cite = self.meta.get("citation_style") or INLINE_FOR_STYLE.get(self.bib_style, "footnotes")
        self.sup = self.meta.get("numeric_superscript", self.loc.get("numeric_superscript", False))
        self.sources = {s["id"]: s for s in spec.get("sources", [])}
        self.source_order: list[str] = []
        self.doc = Document()
        self.n_table = self.n_fig = 0
        self.chart_n = 0
        self.fn_id = 0
        self.footnotes_el = None
        self.docpr_id = 1000
        self.headings: list[tuple[int, str]] = []
        self.report = {"charts": 0, "tables": 0, "footnotes": 0, "warnings": []}
        self._remove_thumbnail()
        self._setup_page()
        self._setup_styles()

    # ---------------------------------------------------------------- setup
    def _remove_thumbnail(self):
        pkg = self.doc.part.package
        for rId, rel in list(pkg.rels.items()):
            if rel.reltype.endswith("/thumbnail"):
                del pkg.rels[rId]

    def _setup_page(self):
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.orientation = WD_ORIENT.PORTRAIT
        if self.gost or self.meta.get("gost_margins"):
            l_, r_, t_, b_ = (3.0, 1.5, 2.0, 2.0)
        else:
            l_, r_, t_, b_ = self.meta.get("margins_cm") or self.loc["margins_cm"]
        sec.left_margin, sec.right_margin = Cm(l_), Cm(r_)
        sec.top_margin, sec.bottom_margin = Cm(t_), Cm(b_)
        self.text_width_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000
        sec.different_first_page_header_footer = True
        # footer: centered page number; header: short title
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(fp, "PAGE", "1", size=self.size - 2)
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run(self.meta.get("header", self.meta.get("title", ""))[:90])
        self._fmt_run(r, size=self.size - 3, color=self.theme["muted"])

    def _setup_styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = self.font
        normal.font.size = Pt(self.size)
        self._east_asia(normal.element, self.font, self.ea_font)
        pf = normal.paragraph_format
        pf.line_spacing = self.spacing
        pf.space_after = Pt(0 if self.gost else 6)
        pf.space_before = Pt(0)
        sizes = {1: self.size + 4, 2: self.size + 2, 3: self.size + 1, 4: self.size}
        for lvl in range(1, 5):
            hs = st[f"Heading {lvl}"]
            hs.font.name = self.hfont
            hs.font.size = Pt(sizes[lvl])
            hs.font.bold = True
            hs.font.italic = False
            hs.font.color.rgb = _rgb("000000") if self.gost else (_rgb(self.accent) if lvl == 1 else _rgb("1A1D23"))
            self._east_asia(hs.element, self.hfont, self.hea_font)
            hpf = hs.paragraph_format
            hpf.space_before = Pt(18 if lvl == 1 else 12)
            hpf.space_after = Pt(8 if lvl == 1 else 6)
            hpf.keep_with_next = True
            hpf.line_spacing = 1.0
            hpf.first_line_indent = Cm(self.indent_cm) if self.gost and lvl > 1 else Cm(0)
            if lvl == 1:
                hpf.page_break_before = bool(self.meta.get("h1_page_break", self.gost))
        cap = st["Caption"]
        cap.font.name = self.font
        cap.font.size = Pt(self.size - 1)
        cap.font.bold = False
        cap.font.italic = False
        cap.font.color.rgb = _rgb("1A1D23")
        cap.paragraph_format.space_after = Pt(4)
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number", "List Bullet 2"):
            s = st[name]
            s.font.name = self.font
            s.font.size = Pt(self.size)
            s.paragraph_format.space_after = Pt(3)
            s.paragraph_format.line_spacing = self.spacing
        tit = st["Title"]
        tit.font.name = self.hfont
        tit.font.size = Pt(26)
        tit.font.color.rgb = _rgb("1A1D23")
        self._east_asia(tit.element, self.hfont, self.hea_font)
        for sty in (cap, st["List Bullet"], st["List Number"], st["List Bullet 2"]):
            self._east_asia(sty.element, self.font, self.ea_font)
        # remove the blue bottom border of the default Title style
        pPr = tit.element.find(qn("w:pPr"))
        if pPr is not None:
            for b in pPr.findall(qn("w:pBdr")):
                pPr.remove(b)

    @staticmethod
    def _east_asia(style_el, font, ea=None):
        rPr = style_el.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
            rPr.insert(0, rFonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rFonts.set(qn(a), font)
        rFonts.set(qn("w:eastAsia"), ea or font)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rFonts.get(qn(a)) is not None:
                del rFonts.attrib[qn(a)]

    def _fmt_run(self, r, size=None, bold=None, italic=None, color=None, font=None):
        if size:
            r.font.size = Pt(size)
        if bold is not None:
            r.font.bold = bold
        if italic is not None:
            r.font.italic = italic
        if color:
            r.font.color.rgb = _rgb(color)
        f = font or self.font
        r.font.name = f
        rPr = r._r.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
            rPr.insert(0, rFonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rFonts.set(qn(a), f)
        rFonts.set(qn("w:eastAsia"), self.ea_map.get(f, self.ea_font if f == self.font else f) if font else self.ea_font)

    def _field(self, paragraph, instr, cached="", size=None):
        r1 = paragraph.add_run()
        r1._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
        r2 = paragraph.add_run()
        r2._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {instr} </w:instrText>'))
        r3 = paragraph.add_run()
        r3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
        r4 = paragraph.add_run(cached)
        r5 = paragraph.add_run()
        r5._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
        for r in (r1, r2, r3, r4, r5):
            self._fmt_run(r, size=size)

    # ---------------------------------------------------------------- footnotes
    def _ensure_footnotes(self):
        if self.footnotes_el is not None:
            return
        xml = (f'<w:footnotes {nsdecls("w", "r")}>'
               '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:separator/></w:r></w:p></w:footnote>'
               '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
               '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote></w:footnotes>')
        self.footnotes_el = etree.fromstring(xml.encode())
        self.footnotes_part = Part(PackURI("/word/footnotes.xml"), CT_FOOTNOTES, b"", self.doc.part.package)
        self.doc.part.relate_to(self.footnotes_part, RT.FOOTNOTES)
        settings = self.doc.settings.element
        if settings.find(qn("w:footnotePr")) is None:
            fp = parse_xml(f'<w:footnotePr {nsdecls("w")}><w:footnote w:id="-1"/><w:footnote w:id="0"/></w:footnotePr>')
            # schema order: footnotePr comes after w:characterSpacingControl-ish elements; insert before compat if present
            insert_before(settings, fp, SETTINGS_TAIL[2:])

    def _add_footnote(self, paragraph, text):
        self._ensure_footnotes()
        self.fn_id += 1
        fid = self.fn_id
        sz = int((self.size - 2) * 2)
        esc = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
        fn = etree.fromstring(
            (f'<w:footnote {nsdecls("w")} w:id="{fid}"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/>'
             f'<w:jc w:val="both"/></w:pPr>'
             f'<w:r><w:rPr><w:rFonts w:ascii="{self.font}" w:hAnsi="{self.font}" w:cs="{self.font}" w:eastAsia="{self.ea_font}"/><w:vertAlign w:val="superscript"/>'
             f'<w:sz w:val="{sz}"/></w:rPr><w:footnoteRef/></w:r>'
             f'<w:r><w:rPr><w:rFonts w:ascii="{self.font}" w:hAnsi="{self.font}" w:cs="{self.font}" w:eastAsia="{self.ea_font}"/><w:sz w:val="{sz}"/></w:rPr>'
             f'<w:t xml:space="preserve"> {esc}</w:t></w:r></w:p></w:footnote>').encode())
        self.footnotes_el.append(fn)
        r = paragraph.add_run()
        self._fmt_run(r)
        rPr = r._r.get_or_add_rPr()
        rPr.append(parse_xml(f'<w:vertAlign {nsdecls("w")} w:val="superscript"/>'))
        r._r.append(parse_xml(f'<w:footnoteReference {nsdecls("w")} w:id="{fid}"/>'))
        self.report["footnotes"] += 1

    def source_text(self, sid):
        """Full reference in the document's bibliography style (gost / apa / chicago / gbt7714).
        Legacy field 'gost' (verbatim string) is honoured for the gost style."""
        s = self.sources.get(sid)
        if not s:
            self.report["warnings"].append(f"unknown source id: {sid}")
            return sid
        if self.bib_style == "gost" and s.get("gost"):
            return s["gost"]
        return format_reference(s, self.bib_style)

    # ---------------------------------------------------------------- inline text
    def _inline(self, paragraph, text, size=None, color=None, bold=None):
        text = re.sub(r"\s+(\[\[(?:src|fn):)", r"\1", text)   # citation runs add their own spacing
        tokens = re.split(r"(\[\[(?:src|fn):[^\]]+\]\]|\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", text)
        for tok in tokens:
            if not tok:
                continue
            m = re.fullmatch(r"\[\[(src|fn):([^\]]+)\]\]", tok)
            if m:
                kind, val = m.groups()
                if kind == "src":
                    ids = [v.strip() for v in val.split(",")]
                    if self.cite == "numeric":
                        nums = []
                        for sid in ids:
                            if sid not in self.source_order:
                                self.source_order.append(sid)
                            nums.append(self.source_order.index(sid) + 1)
                        txt = inline_citation([], "gbt7714", nums)
                        r = paragraph.add_run(txt if self.sup else f" {txt}")
                        self._fmt_run(r, size=size, color=color)
                        if self.sup:
                            r._r.get_or_add_rPr().append(parse_xml(f'<w:vertAlign {nsdecls("w")} w:val="superscript"/>'))
                    elif self.cite == "author_date":
                        for sid in ids:
                            if sid not in self.source_order:
                                self.source_order.append(sid)
                        srcs = [self.sources[x] for x in ids if x in self.sources]
                        for x in ids:
                            if x not in self.sources:
                                self.report["warnings"].append(f"unknown source id: {x}")
                        r = paragraph.add_run(" " + inline_citation(srcs, self.bib_style))
                        self._fmt_run(r, size=size, color=color)
                    else:
                        for sid in ids:
                            if sid not in self.source_order:
                                self.source_order.append(sid)
                        self._add_footnote(paragraph, "; ".join(self.source_text(s) for s in ids))
                else:
                    self._add_footnote(paragraph, val.strip())
                continue
            if tok.startswith("**") and tok.endswith("**"):
                r = paragraph.add_run(tok[2:-2]); self._fmt_run(r, size=size, bold=True, color=color)
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                r = paragraph.add_run(tok[1:-1]); self._fmt_run(r, size=size, italic=True, color=color, bold=bold)
            else:
                r = paragraph.add_run(tok); self._fmt_run(r, size=size, color=color, bold=bold)

    def _para(self, text, style=None, align=None, indent=True, size=None, color=None, space_after=None, bold=None):
        p = self.doc.add_paragraph(style=style)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "left":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent and self.indent_cm:
            p.paragraph_format.first_line_indent = Cm(self.indent_cm)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        self._inline(p, text, size=size, color=color, bold=bold)
        return p

    def _source_line(self, text):
        if not text:
            return
        lab = self.L["source"]
        sep = "：" if self.lang == "zh" else ": "
        txt = text if re.match(r"^(Источник|Источники|Source|Sources|资料来源|来源|数据来源|Nguồn|Sumber)", text) else f"{lab}{sep}{text}"
        p = self._para(txt, align="left", indent=False, size=self.size - 2, color=self.theme["muted"], space_after=10)
        p.paragraph_format.line_spacing = 1.0
        return p

    # ---------------------------------------------------------------- blocks
    def b_title_page(self, b):
        m = self.meta
        for _ in range(3):
            self.doc.add_paragraph()
        if m.get("org"):
            self._para(m["org"], align="left", indent=False, size=self.size, color=self.theme["muted"], space_after=36)
        p = self.doc.add_paragraph(style="Title")
        r = p.add_run(m.get("title", ""))
        self._fmt_run(r, size=26, bold=True, font=self.hfont)
        p.paragraph_format.space_after = Pt(12)
        if m.get("subtitle"):
            self._para(m["subtitle"], align="left", indent=False, size=self.size + 3, color=self.theme["muted"], space_after=48)
        # accent rule
        rule = self.doc.add_paragraph()
        pPr = rule._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="18" w:space="1" w:color="{self.accent}"/></w:pBdr>'))
        if not m.get("date"):
            m["date"] = fmt_date(None, self.lang)
        for key in ("author", "date"):
            if m.get(key):
                self._para(m[key], align="left", indent=False, size=self.size, space_after=2)
        if m.get("watermark"):
            self.doc.add_paragraph()
            self._para(m["watermark"], align="left", indent=False, size=self.size - 1, color=self.theme["negative"], bold=True)
        self.doc.add_paragraph()
        self._para("— " + disclaimer(self.lang) + " —", align="left", indent=False, size=self.size - 2,
                   color=self.theme["muted"])
        self.b_page_break({})

    def b_toc(self, b):
        h = self.doc.add_paragraph()
        r = h.add_run(b.get("title", self.L["toc"]))
        self._fmt_run(r, size=self.size + 4, bold=True, font=self.hfont)
        h.paragraph_format.space_after = Pt(12)
        self._toc_anchor = self.doc.add_paragraph()  # filled after all headings are known
        self._toc_levels = b.get("levels", 2)

    def _fill_toc(self):
        if not hasattr(self, "_toc_anchor"):
            return
        anchor = self._toc_anchor
        p0 = anchor
        self._field_begin(p0, f'TOC \\o "1-{self._toc_levels}" \\h \\z \\u')
        prev = p0
        entries = [(lvl, t) for lvl, t in self.headings if lvl <= self._toc_levels]
        for lvl, text in entries:
            np_ = parse_xml(f'<w:p {nsdecls("w")}/>')
            prev._p.addnext(np_) if hasattr(prev, "_p") else prev.addnext(np_)
            from docx.text.paragraph import Paragraph
            para = Paragraph(np_, anchor._parent)
            para.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
            para.paragraph_format.space_after = Pt(2)
            r = para.add_run(text)
            self._fmt_run(r, size=self.size, bold=(lvl == 1))
            prev = para
        endp = parse_xml(f'<w:p {nsdecls("w")}/>')
        prev._p.addnext(endp)
        from docx.text.paragraph import Paragraph
        ep = Paragraph(endp, anchor._parent)
        r = ep.add_run()
        r._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
        # ask Word to refresh fields (page numbers) on open
        st = self.doc.settings.element
        if st.find(qn("w:updateFields")) is None:
            insert_before(st, parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'), SETTINGS_TAIL)

    def _field_begin(self, paragraph, instr):
        r1 = paragraph.add_run()
        r1._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin" w:dirty="true"/>'))
        r2 = paragraph.add_run()
        r2._r.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {instr} </w:instrText>'))
        r3 = paragraph.add_run()
        r3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))

    def b_page_break(self, b):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    def b_heading(self, b):
        lvl = int(b.get("level", 1))
        p = self.doc.add_heading(level=lvl)
        self._inline(p, b["text"], bold=True)
        for r in p.runs:
            r.font.color.rgb = None if self.gost else r.font.color.rgb
        self.headings.append((lvl, re.sub(r"\[\[[^\]]+\]\]", "", b["text"]).replace("**", "")))

    def b_paragraph(self, b):
        self._para(b["text"], align=b.get("align"), indent=b.get("indent", True))

    def b_bullets(self, b):
        for it in b["items"]:
            text, lvl = (it.get("text", ""), int(it.get("level", 0))) if isinstance(it, dict) else (str(it), 0)
            p = self.doc.add_paragraph(style="List Bullet 2" if lvl else "List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self._inline(p, text)

    def b_numbered(self, b):
        num_id = self._new_num_restart()
        for it in b["items"]:
            p = self.doc.add_paragraph(style="List Number")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_num(p, num_id)
            self._inline(p, it if isinstance(it, str) else it.get("text", ""))

    def _new_num_restart(self):
        numbering = self.doc.part.numbering_part.element
        style = self.doc.styles["List Number"].element
        base_num = style.find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
        abs_id = None
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == base_num:
                abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        new_id = max(ids) + 1
        numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{new_id}"><w:abstractNumId w:val="{abs_id}"/>'
            f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
        return new_id

    def _caption(self, kind, text):
        if kind == "table":
            self.n_table += 1
            label = f"{self.L['table']}{self.loc['caption_num_sep']}{self.n_table}"
        else:
            self.n_fig += 1
            label = f"{self.L['figure']}{self.loc['caption_num_sep']}{self.n_fig}"
        p = self.doc.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if kind == "table" else WD_ALIGN_PARAGRAPH.CENTER
        if kind == "table":
            p.paragraph_format.keep_with_next = True
            if self.lang == "zh":   # GB/T convention: table and figure captions centred
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}{self.loc['caption_sep']}")
        self._fmt_run(r, size=self.size - 1, bold=True)
        self._inline(p, text or "", size=self.size - 1)
        return p

    def _fix_widths(self, t, widths_cm):
        """Set tblGrid/gridCol + fixed layout so Word *and* LibreOffice honour column widths."""
        for i, w in enumerate(widths_cm):
            t.columns[i].width = Cm(w)
        tblPr = t._tbl.tblPr
        lay = tblPr.find(qn("w:tblLayout"))
        if lay is None:
            insert_before(tblPr, parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'), ["tblCellMar", "tblLook", "tblCaption", "tblDescription"])
        tw = tblPr.find(qn("w:tblW"))
        if tw is not None:
            tw.set(qn("w:type"), "dxa")
            tw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))

    def _cell_shade(self, cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color}"/>'))

    def _table_borders(self, table, color="BFBFBF", inside=True):
        tblPr = table._tbl.tblPr
        ins = (f'<w:insideH w:val="single" w:sz="4" w:color="{color}"/>' if inside else "")
        insert_before(tblPr, parse_xml(
            f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="8" w:color="1A1D23"/>'
            f'<w:left w:val="nil"/><w:bottom w:val="single" w:sz="8" w:color="1A1D23"/><w:right w:val="nil"/>'
            f'{ins}<w:insideV w:val="nil"/></w:tblBorders>'), TBLPR_AFTER_BORDERS)

    def b_table(self, b):
        if b.get("caption"):
            self._caption("table", b["caption"])
        cols, rows = b["columns"], b["rows"]
        t = self.doc.add_table(rows=len(rows) + 1, cols=len(cols))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._table_borders(t)
        widths = b.get("col_widths_cm")
        if not widths:
            first = 0.34 if len(cols) > 2 else 0.5
            rest = (1 - first) / max(len(cols) - 1, 1)
            widths = [self.text_width_cm * first] + [self.text_width_cm * rest] * (len(cols) - 1)
        num_re = re.compile(r"^[\s+\-−–]?(?:Rp|RM|\$)?[\d\s\u00a0.,]+\s*(%|x|×|млрд|млн|тыс\.?|₽|\$|мес\.?|лет|ч|дн\.?|п\.п\.|pp|bn|mn|万元|亿元|万|亿|元|倍|个百分点|个月|年|天|小时|nghìn tỷ|tỷ|triệu|₫|lần|điểm %|tháng|triliun|miliar|juta|bilion|ribu|bulan|hari|jam|giờ|ngày)?[\s\u00a0₽$¥₫]*$")
        fs = b.get("font_size", self.size - 1)
        tot_w = sum(widths)
        if b.get("fit_width", True) or tot_w > self.text_width_cm + 0.01:
            widths = [w * self.text_width_cm / tot_w for w in widths]
        self._fix_widths(t, widths)
        num_cols = [c > 0 and all(bool(num_re.match(str(row[c]).replace("**", ""))) or str(row[c]) in ("", "—", "-")
                                  for row in rows) for c in range(len(cols))]
        for r_i in range(len(rows) + 1):
            for c_i in range(len(cols)):
                cell = t.cell(r_i, c_i)
                cell.width = Cm(widths[c_i])
                val = cols[c_i] if r_i == 0 else rows[r_i - 1][c_i]
                p = cell.paragraphs[0]
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                if len(rows) <= 15 and b.get("keep_together", True):
                    p.paragraph_format.keep_with_next = True
                p.paragraph_format.line_spacing = 1.0
                is_num = num_cols[c_i]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if is_num else WD_ALIGN_PARAGRAPH.LEFT
                self._inline(p, str(val), size=fs, bold=(r_i == 0) or (b.get("bold_first_col") and c_i == 0))
                if r_i == 0:
                    self._cell_shade(cell, "F2F2F2")
                elif b.get("highlight_row") is not None and r_i - 1 == b["highlight_row"]:
                    self._cell_shade(cell, "E8F3F1")
        # repeat header row
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        self.report["tables"] += 1
        self._source_line(b.get("source"))
        if not b.get("source"):
            self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def b_glossary(self, b):
        """Plain-language definitions of the metrics used in the report (tools/glossary.py)."""
        ids = b.get("ids") or []
        if not ids:
            return
        self.b_heading({"level": 1, "text": glossary.label("title", self.lang)})
        self._para(glossary.label("intro", self.lang))
        rows = [[e["term"], e["text"]] for e in glossary.entries(ids, self.lang)]
        self.b_table({"columns": [glossary.label("col_term", self.lang), glossary.label("col_def", self.lang)],
                      "rows": rows, "col_widths_cm": [4.2, 12.3], "bold_first_col": True, "keep_together": False,
                      "source": glossary.label("source", self.lang)})
        tblPr = self.doc.tables[-1]._tbl.tblPr
        tblPr.append(parse_xml(f'<w:tblCaption {nsdecls("w")} w:val="{glossary.GLOSSARY_TABLE_CAPTION}"/>'))
        self.report["glossary"] = ids

    def b_key_figures(self, b):
        items = b["items"][:4]
        t = self.doc.add_table(rows=2, cols=len(items))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        insert_before(tblPr, parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="12" w:color="{self.accent}"/>'
                                       f'<w:left w:val="nil"/><w:bottom w:val="nil"/><w:right w:val="nil"/><w:insideH w:val="nil"/>'
                                       f'<w:insideV w:val="nil"/></w:tblBorders>'), TBLPR_AFTER_BORDERS)
        self._fix_widths(t, [self.text_width_cm / len(items)] * len(items))
        for i, it in enumerate(items):
            c0, c1 = t.cell(0, i), t.cell(1, i)
            for c in (c0, c1):
                c.width = Cm(self.text_width_cm / len(items))
            p = c0.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(6)
            r = p.add_run(str(it["value"]))
            self._fmt_run(r, size=self.size + 10, bold=True, color=self.accent, font=self.hfont)
            p2 = c1.paragraphs[0]
            p2.paragraph_format.first_line_indent = Cm(0)
            p2.paragraph_format.line_spacing = 1.0
            self._inline(p2, it.get("label", ""), size=self.size - 1)
        self._source_line(b.get("source"))

    def b_callout(self, b):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        cell.width = Cm(self.text_width_cm)
        self._fix_widths(t, [self.text_width_cm])
        self._cell_shade(cell, "F3F4F6")
        tblPr = t._tbl.tblPr
        insert_before(tblPr, parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="nil"/><w:left w:val="single" w:sz="24" w:color="{self.accent}"/>'
                                       f'<w:bottom w:val="nil"/><w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>'), TBLPR_AFTER_BORDERS)
        p = cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Cm(0)
        if b.get("title"):
            r = p.add_run(b["title"] + " ")
            self._fmt_run(r, bold=True, color=self.accent)
        self._inline(p, b.get("text", ""))
        for extra in b.get("bullets", []):
            q = cell.add_paragraph(style="List Bullet")
            self._inline(q, extra)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def b_scorecard(self, b):
        crit = b["criteria"]
        rows = []
        dec = lambda v: fmt_number(v, 1, self.lang)
        tot_w = sum(c["weight"] for c in crit)
        for c in crit:
            contrib = c["weight"] * c["score"] / 5
            rows.append([c["name"], dec(c["weight"]), dec(c["score"]), dec(contrib), c.get("note", "")])
        total = b.get("total", sum(c["weight"] * c["score"] / 5 for c in crit))
        rows.append([f"**{self.L['total']}**", dec(tot_w), "", f"**{dec(total)}**", b.get("total_note", "")])
        has_note = any(c.get("note") for c in crit)
        cols = [self.L["criterion"], self.L["weight"], self.L["score"] + " (0–5)", self.L["contribution"]]
        widths = [6.0, 2.0, 2.2, 2.2]
        if has_note:
            cols.append(self.L["comment"])
            rem = self.text_width_cm - sum(widths)
            widths.append(max(rem, 3.0))
        else:
            rows = [r[:4] for r in rows]
            widths = [self.text_width_cm - 6.4, 2.0, 2.2, 2.2]
        self.b_table({"caption": b.get("caption"), "columns": cols, "rows": rows, "source": b.get("source"),
                      "col_widths_cm": widths, "font_size": self.size - 2})

    def b_chart(self, b):
        cspec = b["chart"]
        xl_type, cdata, norm = prepare(cspec)
        from pptx.chart.chart import Chart
        from pptx.oxml import parse_xml as pptx_parse
        cs = pptx_parse(cdata.xml_bytes(xl_type))
        chart = Chart(cs, None)
        th = dict(self.theme)
        th["font_body"] = self.meta.get("chart_font", self.font)
        th["font_ea"] = self.meta.get("chart_font_ea", self.ea_font if self.ea_font != self.font else None)
        th["lang"] = self.lang
        th["palette"] = cspec.get("colors") or [self.accent] + self.theme["palette"][1:]
        style_chart(chart, norm, th, base_font_pt=b.get("font_pt", 9))
        fix_axis_ids(cs)
        set_chart_lang(cs, self.lang)
        if th.get("font_ea"):
            add_ea_fonts(cs, th["font_ea"], None, th["font_body"])
        self.chart_n += 1
        n = self.chart_n
        pkg = self.doc.part.package
        xlsx_part = Part(PackURI(f"/word/embeddings/Microsoft_Excel_Worksheet{n}.xlsx"), CT_XLSX, cdata.xlsx_blob, pkg)
        chart_part = Part(PackURI(f"/word/charts/chart{n}.xml"), CT_CHART, b"", pkg)
        x_rid = chart_part.relate_to(xlsx_part, RT_PACKAGE)
        ext = cs.get_or_add_externalData()
        ext.rId = x_rid
        if ext.find(qn("c:autoUpdate")) is None:
            ext.append(etree.fromstring('<c:autoUpdate xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" val="0"/>'))
        chart_part._blob = etree.tostring(cs, xml_declaration=True, encoding="UTF-8", standalone=True)
        c_rid = self.doc.part.relate_to(chart_part, RT.CHART)
        w_cm = min(b.get("width_cm", self.text_width_cm), self.text_width_cm)
        h_cm = b.get("height_cm", 8.0)
        cx, cy = int(Cm(w_cm)), int(Cm(h_cm))
        self.docpr_id += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        inline = parse_xml(
            f'<wp:inline {nsdecls("wp", "a", "r")} distT="0" distB="0" distL="0" distR="0">'
            f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
            f'<wp:docPr id="{self.docpr_id}" name="Chart {n}" descr="{(b.get("caption") or "chart").replace(chr(34), "")}"/>'
            f'<wp:cNvGraphicFramePr/>'
            f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            f'<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" r:id="{c_rid}"/>'
            f'</a:graphicData></a:graphic></wp:inline>')
        run = p.add_run()
        drawing = parse_xml(f'<w:drawing {nsdecls("w")}/>')
        drawing.append(inline)
        run._r.append(drawing)
        if b.get("caption"):
            self._caption("figure", b["caption"])
        self._source_line(b.get("source"))
        self.report["charts"] += 1

    def b_sources(self, b):
        title = b.get("title", self.L["references"])
        h = self.doc.add_heading(level=1)
        self._inline(h, title)
        self.headings.append((1, title))
        uncited = [k for k in self.sources if k not in self.source_order]
        if self.cite == "author_date":   # APA / Chicago: alphabetical, hanging indent, unnumbered
            order = sorted(self.source_order + uncited, key=lambda k: sort_key(self.sources.get(k, {"title": k}), self.bib_style))
            for sid in order:
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Cm(1.27)
                p.paragraph_format.first_line_indent = Cm(-1.27)
                self._fmt_run(p.add_run(self.source_text(sid)), size=self.size - 1)
            return
        order = self.source_order + uncited
        if self.cite == "numeric":       # GB/T 7714 顺序编码制: [1] ... in citation order
            for i, sid in enumerate(order, 1):
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.9)
                self._fmt_run(p.add_run(f"[{i}]\t{self.source_text(sid)}"), size=self.size - 1)
            return
        num_id = self._new_num_restart()  # ГОСТ: numbered list in order of first citation
        for sid in order:
            p = self.doc.add_paragraph(style="List Number")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_num(p, num_id)
            r = p.add_run(self.source_text(sid))
            self._fmt_run(r, size=self.size - 1)

    # ---------------------------------------------------------------- build
    def build(self):
        for b in self.spec["blocks"]:
            fn = getattr(self, f"b_{b['type']}", None)
            if fn is None:
                self.report["warnings"].append(f"unknown block type {b['type']}")
                continue
            fn(b)
        self.doc.add_paragraph()
        h = self.doc.add_paragraph()
        r = h.add_run(disclaimer_title(self.lang))
        self._fmt_run(r, size=self.size + 2, bold=True, font=self.hfont)
        h.paragraph_format.space_after = Pt(6)
        self._para(disclaimer(self.lang), align="left", indent=False, size=self.size - 1, color=self.theme["muted"])
        self._fill_toc()
        if self.footnotes_el is not None:
            self.footnotes_part._blob = etree.tostring(self.footnotes_el, xml_declaration=True, encoding="UTF-8",
                                                       standalone=True)
        cp = self.doc.core_properties
        cp.title = self.meta.get("title", "")
        cp.author = self.meta.get("author", "")
        # never leak library-template defaults (template author, "generated by ...") into metadata
        cp.last_modified_by = self.meta.get("author", "")
        cp.comments = self.meta.get("comments", "")
        cp.keywords = self.meta.get("keywords", "")
        cp.subject = self.meta.get("subject", "")
        cp.revision = 1
        cp.language = self.loc["lang_tag"]
        # document language for proofing and East Asian line breaking
        rpr = self.doc.styles["Normal"].element.get_or_add_rPr()
        ea = ' w:eastAsia="zh-CN"' if self.lang == "zh" else ""
        rpr.append(parse_xml(f'<w:lang {nsdecls("w")} w:val="{"en-US" if self.lang == "zh" else cp.language}"{ea}/>'))
        return self.doc


def build_docx(spec: dict | str, out_path: str, facts: dict | str | None = None, font: str | None = None,
               markdown: str | None = None, meta: dict | str | None = None, locale: str | None = None,
               citation_style: str | None = None) -> dict:
    if markdown:
        with open(markdown, encoding="utf-8") as fh:
            spec = {"meta": load_json(meta) if meta else {}, "blocks": markdown_to_blocks(fh.read())}
    else:
        spec = load_json(spec)
    spec = copy.deepcopy(spec)
    if font:
        spec.setdefault("meta", {})["font"] = font
    if locale:
        spec.setdefault("meta", {})["locale"] = locale
    if citation_style:
        spec.setdefault("meta", {})["bibliography_style"] = citation_style
        spec["meta"]["citation_style"] = INLINE_FOR_STYLE[citation_style]
    fsrc = facts or spec.get("meta", {}).get("facts")
    if fsrc:
        spec = resolve_placeholders(spec, load_json(fsrc), {"locale": locale_of(spec.get("meta")),
                                                            "currency": spec["meta"].get("currency")})
    if spec.get("meta", {}).get("glossary", True) and not any(b.get("type") == "glossary" and b.get("ids")
                                                              for b in spec.get("blocks", [])):
        # pass 1: build, read the visible text exactly as qa_check does, detect the metrics shown
        import io
        buf = io.BytesIO()
        ReportBuilder(copy.deepcopy(spec)).build().save(buf)
        ids = glossary.detect(glossary.docx_parts(buf.getvalue())["text"], locale_of(spec.get("meta")))
        if ids:
            spec = glossary.insert_report_block(spec, ids)
    rb = ReportBuilder(spec)
    doc = rb.build()
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    clean_package(out_path, rb.lang)   # drop unrelated template theme fonts / languages
    rb.report["output"] = os.path.basename(out_path)
    rb.report["headings"] = len(rb.headings)
    return rb.report


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", "-i")
    ap.add_argument("--markdown", "-m")
    ap.add_argument("--meta")
    ap.add_argument("--out", "-o", required=True)
    ap.add_argument("--facts")
    ap.add_argument("--font")
    ap.add_argument("--locale", "-l", choices=["ru", "en", "zh", "vi", "id", "ms"], help="overrides meta.locale")
    ap.add_argument("--citation-style", choices=["gost", "apa", "chicago", "gbt7714"],
                    help="bibliography style; in-text mode follows (footnotes / author-date / numeric)")
    a = ap.parse_args(argv)
    if not a.input and not a.markdown:
        ap.error("--input or --markdown required")
    print(dump_json(build_docx(a.input, a.out, a.facts, a.font, a.markdown, a.meta, a.locale, a.citation_style)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
