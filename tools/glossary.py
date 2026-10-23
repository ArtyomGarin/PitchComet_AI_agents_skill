#!/usr/bin/env python3
"""glossary.py — plain-language "What the numbers mean" slide(s) and report section.

Definitions live in config/glossary.json (ru, en, zh, vi, id, ms). The glossary lists only the
metrics a document actually shows: build_pptx / build_docx build the document once, read its
visible text with the same extractor qa_check uses, detect the metrics and build it again with
the glossary inserted (slide 2 of a deck, the section before the first heading of a report).

CLI:
  python glossary.py list --locale ru                 # every term with its definition
  python glossary.py detect --pptx deck.pptx          # metrics shown in a deck
  python glossary.py detect --text "LTV/CAC 5x, CAGR 18%" --locale en
"""
from __future__ import annotations

import argparse
import io
import json
import math
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
GLOSSARY_FILE = os.path.join(HERE, "..", "config", "glossary.json")
LOCALES = ("ru", "en", "zh", "vi", "id", "ms")
_CACHE: dict = {}

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
GLOSSARY_TABLE_CAPTION = "pitchcomet-glossary"
TERM_SHAPE = "GlossaryTerm"
DEF_SHAPE = "GlossaryDef"


def load(path: str | None = None) -> dict:
    path = path or GLOSSARY_FILE
    if path not in _CACHE:
        with open(path, encoding="utf-8") as fh:
            _CACHE[path] = json.load(fh)
    return _CACHE[path]


def _loc(locale: str | None) -> str:
    l = (locale or "en").lower().split("-")[0].split("_")[0]
    return l if l in LOCALES else "en"


def _patterns(m: dict, locale: str) -> list[re.Pattern]:
    key = ("pat", m["id"], locale)
    if key in _CACHE:
        return _CACHE[key]
    pats = [re.compile(r"(?<![A-Za-z0-9])" + re.escape(a) + r"(?![A-Za-z])") for a in m.get("abbr", [])]
    words = m.get("words", {})
    for k in {"all", "en", locale}:
        for w in words.get(k, []):
            pats.append(re.compile(w, re.I))
    _CACHE[key] = pats
    return pats


def order() -> list[str]:
    return [m["id"] for m in load()["metrics"]]


def detect(text: str | list, locale: str = "en") -> list[str]:
    """Metric ids shown in text, in glossary order, with implied metrics added."""
    loc = _loc(locale)
    if isinstance(text, (list, tuple)):
        text = "\n".join(str(t) for t in text)
    found = set()
    by_id = {m["id"]: m for m in load()["metrics"]}
    for m in load()["metrics"]:
        if any(p.search(text) for p in _patterns(m, loc)):
            found.add(m["id"])
    for i in list(found):
        found.update(by_id[i].get("implies", []))
    return [i for i in order() if i in found]


def entries(ids: list[str], locale: str) -> list[dict]:
    loc = _loc(locale)
    by_id = {m["id"]: m for m in load()["metrics"]}
    return [{"id": i, "term": by_id[i]["term"][loc], "text": by_id[i]["def"][loc]} for i in ids if i in by_id]


def label(key: str, locale: str) -> str:
    return load()["labels"][key][_loc(locale)]


def titles(locale: str | None = None) -> set[str]:
    t = load()["labels"]["title"]
    return {t[_loc(locale)]} if locale else set(t.values())


def term_to_id(term: str, locale: str | None = None) -> str | None:
    term = (term or "").strip()
    for m in load()["metrics"]:
        vals = [m["term"][_loc(locale)]] if locale else m["term"].values()
        if term in vals:
            return m["id"]
    return None


# ----------------------------------------------------------------------------- deck
def slide_specs(ids: list[str], locale: str, per_slide: int | None = None) -> list[dict]:
    items = entries(ids, locale)
    if not items:
        return []
    per = per_slide or load().get("per_slide", 8)
    k = math.ceil(len(items) / per)
    size = math.ceil(len(items) / k)
    out = []
    for n in range(k):
        chunk = items[n * size:(n + 1) * size]
        title = label("title", locale) + ("" if n == 0 else " " + label("continued", locale))
        out.append({"layout": "glossary", "title": title, "items": chunk, "source": label("source", locale),
                    "notes": label("notes", locale), "_glossary": True})
    return out


def insert_slides(spec: dict, ids: list[str], locale: str) -> dict:
    slides = [s for s in spec.get("slides", []) if not s.get("_glossary")]
    gl = slide_specs(ids, locale)
    pos = 1 if slides and slides[0].get("layout") == "title" else 0
    spec["slides"] = slides[:pos] + gl + slides[pos:]
    return spec


def _chart_xml_text(xml: bytes) -> str:
    vals = re.findall(rb"<(?:a:t|c:v)>([^<]*)</(?:a:t|c:v)>", xml)
    return " ".join(v.decode("utf-8", "replace") for v in vals if re.search(rb"[^\d\s.,\-E+]", v))


def pptx_slide_texts(prs) -> list[tuple[int, bool, str]]:
    """[(slide number, is_glossary, visible text incl. tables and native charts)] — notes excluded."""
    out = []
    for idx, slide in enumerate(prs.slides, 1):
        parts, is_gl = [], False

        def walk(shapes):
            nonlocal is_gl
            for shp in shapes:
                nm = shp.name or ""
                if nm.startswith(TERM_SHAPE):
                    is_gl = True
                if getattr(shp, "shape_type", None) is not None and int(shp.shape_type) == 6:   # group
                    walk(shp.shapes)
                    continue
                if getattr(shp, "has_chart", False) and shp.has_chart:
                    try:
                        parts.append(_chart_xml_text(shp.chart.part.blob))
                    except Exception:
                        pass
                    continue
                if getattr(shp, "has_table", False) and shp.has_table:
                    parts.extend(c.text for r in shp.table.rows for c in r.cells)
                    continue
                if shp.has_text_frame and nm not in ("PageNumber", "Watermark"):
                    parts.append(shp.text_frame.text)
        walk(slide.shapes)
        out.append((idx, is_gl, "\n".join(parts)))
    return out


def pptx_glossary_ids(prs) -> tuple[list[int], list[str]]:
    slides, ids = [], []
    for idx, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            nm = shp.name or ""
            if nm.startswith(TERM_SHAPE + ":"):
                if idx not in slides:
                    slides.append(idx)
                ids.append(nm.split(":", 1)[1])
    return slides, ids


def docx_parts(data: bytes) -> dict:
    """Visible body text of a .docx split into the glossary section and the rest (+ chart text)."""
    from lxml import etree
    z = zipfile.ZipFile(io.BytesIO(data))
    doc = etree.fromstring(z.read("word/document.xml"))
    body = doc.find(f"{{{W_NS}}}body")
    elems = list(body)

    def ptext(e):
        return "".join(t.text or "" for t in e.iter(f"{{{W_NS}}}t"))

    gl_idx, gl_terms = None, []
    for i, el in enumerate(elems):
        if etree.QName(el).localname == "tbl":
            cap = el.find(f"{{{W_NS}}}tblPr/{{{W_NS}}}tblCaption")
            if cap is not None and cap.get(f"{{{W_NS}}}val") == GLOSSARY_TABLE_CAPTION:
                gl_idx = i
                for r_i, tr in enumerate(el.findall(f"{{{W_NS}}}tr")):
                    tcs = tr.findall(f"{{{W_NS}}}tc")
                    if r_i and tcs:
                        gl_terms.append(ptext(tcs[0]).strip())
                break
    start = gl_idx
    if gl_idx is not None:   # the section starts at the heading just before the table
        for j in range(gl_idx - 1, max(-1, gl_idx - 5), -1):
            sty = elems[j].find(f"{{{W_NS}}}pPr/{{{W_NS}}}pStyle")
            if sty is not None and sty.get(f"{{{W_NS}}}val", "").lower().startswith("heading"):
                start = j
                break
    rest = [ptext(e) for i, e in enumerate(elems) if gl_idx is None or not (start <= i <= gl_idx + 1)]
    charts = " ".join(_chart_xml_text(z.read(n)) for n in z.namelist() if n.startswith("word/charts/chart") and n.endswith(".xml"))
    n_before = sum(1 for e in elems[:start or 0]
                   if (s := e.find(f"{{{W_NS}}}pPr/{{{W_NS}}}pStyle")) is not None and s.get(f"{{{W_NS}}}val", "").lower().startswith("heading"))
    return {"has_glossary": gl_idx is not None, "terms": gl_terms, "text": "\n".join(rest) + "\n" + charts,
            "headings_before": n_before}


def report_block(ids: list[str], locale: str) -> dict:
    return {"type": "glossary", "ids": ids}


def insert_report_block(spec: dict, ids: list[str]) -> dict:
    blocks = [b for b in spec.get("blocks", []) if b.get("type") != "glossary" or b.get("ids") is None or b.get("_manual")]
    blocks = [b for b in blocks if b.get("type") != "glossary"]
    pos = next((i for i, b in enumerate(blocks) if b.get("type") == "heading"), None)
    if pos is None:
        pos = next((i + 1 for i, b in enumerate(blocks) if b.get("type") in ("title_page", "toc", "page_break")), 0)
        while pos < len(blocks) and blocks[pos].get("type") in ("toc", "page_break"):
            pos += 1
    spec["blocks"] = blocks[:pos] + [report_block(ids, None)] + blocks[pos:]
    return spec


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["list", "detect"])
    ap.add_argument("--locale", "-l", default="en", choices=list(LOCALES))
    ap.add_argument("--pptx")
    ap.add_argument("--docx")
    ap.add_argument("--text")
    a = ap.parse_args(argv)
    if a.cmd == "list":
        out = entries(order(), a.locale)
    else:
        if a.pptx:
            from pptx import Presentation
            txt = [t for _, g, t in pptx_slide_texts(Presentation(a.pptx)) if not g]
        elif a.docx:
            txt = docx_parts(open(a.docx, "rb").read())["text"]
        else:
            txt = a.text or sys.stdin.read()
        out = detect(txt, a.locale)
    print(json.dumps(out, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
