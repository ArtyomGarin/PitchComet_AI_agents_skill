"""Unit + integration tests (stdlib unittest; no pytest needed).

Run from the repository root:  python -m unittest discover -s tools/tests -v
All tests write to temporary directories; the repository is never modified.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
TOOLS = os.path.dirname(HERE)
ROOT = os.path.dirname(TOOLS)
sys.path.insert(0, TOOLS)

import metrics_calc as M  # noqa: E402
from build_docx import build_docx  # noqa: E402
from build_pptx import build_pptx  # noqa: E402
from citations import format_reference, inline_citation  # noqa: E402
from financial_model import build as fin_build  # noqa: E402
from fmt_utils import fmt_money, fmt_number, fmt_pct, format_value, load_json, resolve_placeholders  # noqa: E402
from locales import LOCALES, MARKETS, fmt_date, norm_market  # noqa: E402
from market_sizing import bottom_up, size_market, top_down, triangulate  # noqa: E402
from qa_check import run_qa  # noqa: E402
from scoring import allocate_slides, load_model, score, weights_for  # noqa: E402
from textfit import lines_needed, text_units  # noqa: E402

LOCS = ["ru", "en", "zh", "vi", "id", "ms"]
NB = "\u00a0"


def EX(loc: str) -> str:
    return os.path.join(ROOT, "examples", f"example_{loc}")


def v(res):
    return res["value"]


def run_cli(*args):
    return subprocess.run([sys.executable, *args], capture_output=True, text=True, check=True,
                          stdin=subprocess.DEVNULL)


def xml_text(path: str, prefix: str) -> str:
    with zipfile.ZipFile(path) as z:
        return "".join(z.read(n).decode("utf8", "ignore") for n in z.namelist()
                       if n.startswith(prefix) and n.endswith(".xml"))


# ---------------------------------------------------------------- metrics
class TestMetrics(unittest.TestCase):
    def test_cagr(self):
        self.assertAlmostEqual(v(M.cagr(100, 200, 5)), 0.148698, places=5)

    def test_irr_npv(self):
        self.assertAlmostEqual(v(M.irr([-100, 10, 10, 110])), 0.10, places=6)
        self.assertAlmostEqual(v(M.npv(0.10, [-100, 10, 10, 110])), 0.0, places=6)

    def test_moic_irr_from_moic(self):
        self.assertAlmostEqual(v(M.moic(300, 100)), 3.0)
        self.assertAlmostEqual(v(M.irr_from_moic(3.0, 5)), 3 ** 0.2 - 1, places=6)

    def test_saas(self):
        self.assertAlmostEqual(v(M.nrr(100, 20, 5, 5)), 1.10)
        self.assertAlmostEqual(v(M.grr(100, 5, 5)), 0.90)
        self.assertAlmostEqual(v(M.arr_from_mrr(10)), 120)
        self.assertAlmostEqual(v(M.rule_of_40(0.25, 0.20)), 0.45)   # neutral illustrative inputs
        self.assertEqual(M.rule_of_40(0.25, 0.20)["interpretation"], "meets Rule of 40")
        self.assertAlmostEqual(v(M.rule_of_40(0.10, 0.15)), 0.25)
        self.assertAlmostEqual(v(M.burn_multiple(15, 10)), 1.5)
        self.assertAlmostEqual(v(M.cac(1000, 10)), 100)
        self.assertAlmostEqual(v(M.ltv_cac(300, 100)), 3.0)
        self.assertAlmostEqual(v(M.runway(120, 10)), 12)

    def test_real_rate_and_deflate(self):
        self.assertAlmostEqual(v(M.real_rate(0.18, 0.06)), 1.18 / 1.06 - 1, places=9)
        d = M.deflate([100, 106], 0.06)
        self.assertAlmostEqual(d["series"][1], 100.0, places=6)

    def test_wright_break_even_trl_accuracy(self):
        self.assertAlmostEqual(v(M.wright_law(100, 2, 0.2)), 80, places=6)
        self.assertAlmostEqual(v(M.break_even(1000, 50, 30)), 50)
        self.assertEqual(M.trl_assess({f"trl{i}": True for i in range(1, 7)})["value"], 6)
        self.assertEqual(M.trl_assess({"trl1": True, "trl3": True})["value"], 1)
        self.assertAlmostEqual(M.accuracy_vs_baseline(1880, 1640, 2000)["value"], 0.12, places=6)

    def test_every_metric_has_contract(self):
        r = M.cagr(1, 2, 1)
        for k in ("metric", "value", "unit", "formula", "interpretation", "inputs"):
            self.assertIn(k, r)
        self.assertGreaterEqual(len(M.METRICS), 40)

    def test_batch_survives_errors(self):
        out = M.run_batch({"calcs": [{"id": "ok", "fn": "cagr", "args": {"start": 1, "end": 2, "years": 1}},
                                     {"id": "bad", "fn": "cagr", "args": {"start": 0}}]})
        self.assertAlmostEqual(out["results"]["ok"]["value"], 1.0)
        self.assertIn("error", out["results"]["bad"])

    def test_cli_locale_display(self):
        for loc, expect in [("vi", "45,0%"), ("ms", "45.0%"), ("zh", "45.0%")]:
            out = run_cli(os.path.join(TOOLS, "metrics_calc.py"), "rule_of_40", "--json",
                          '{"revenue_growth":0.25,"profit_margin":0.20}', "--locale", loc)
            self.assertEqual(json.loads(out.stdout)["display"], expect, loc)


# ---------------------------------------------------------------- locales and formatting
class TestLocaleFormatting(unittest.TestCase):
    MONEY = {  # locale: (currency, big amount, expected, small amount, expected)
        "ru": ("RUB", 3e9, f"3,0{NB}млрд{NB}₽", 36e6, f"36{NB}млн{NB}₽"),
        "en": ("USD", 3e9, "$3.0B", 36e6, "$36M"),
        "zh": ("CNY", 3e9, "30.0亿元", 36e6, "3,600万元"),
        "vi": ("VND", 3e12, f"3,0{NB}nghìn tỷ{NB}₫", 36e9, f"36{NB}tỷ{NB}₫"),
        "id": ("IDR", 3e12, f"Rp3,0{NB}triliun", 36e9, f"Rp36{NB}miliar"),
        "ms": ("MYR", 3e9, f"RM3.0{NB}bilion", 36e6, f"RM36{NB}juta"),
    }

    def test_all_locales_registered(self):
        self.assertEqual(sorted(LOCALES), sorted(LOCS))
        self.assertEqual(sorted(MARKETS), sorted(["ru", "en", "zh", "vn", "id", "my"]))

    def test_money_per_locale(self):
        for loc, (cur, big, eb, small, es) in self.MONEY.items():
            self.assertEqual(fmt_money(big, cur, loc, "auto", 1), eb, loc)
            self.assertEqual(fmt_money(small, cur, loc, "auto", 0), es, loc)

    def test_negative_money_sign_first(self):
        self.assertEqual(fmt_money(-1e6, "USD", "en", "auto", 0), "-$1M")
        self.assertTrue(fmt_money(-36e9, "IDR", "id", "auto", 0).startswith("-Rp"))
        self.assertTrue(fmt_money(-36e6, "MYR", "ms", "auto", 0).startswith("-RM"))

    def test_separators(self):
        self.assertEqual(fmt_number(12345.6, 1, "vi"), "12.345,6")   # dot thousands, comma decimal
        self.assertEqual(fmt_number(12345.6, 1, "id"), "12.345,6")
        self.assertEqual(fmt_number(12345.6, 1, "ms"), "12,345.6")   # comma thousands, dot decimal
        self.assertEqual(fmt_number(12345.6, 1, "en"), "12,345.6")
        self.assertEqual(fmt_number(12345.6, 1, "ru"), f"12{NB}345,6")
        self.assertEqual(fmt_pct(0.185, 1, "vi"), "18,5%")
        self.assertEqual(fmt_pct(0.185, 1, "ms"), "18.5%")

    def test_points_and_multiples(self):
        self.assertEqual(format_value(0.13, "pp0", {"locale": "vi"}), f"13{NB}điểm %")
        self.assertEqual(format_value(0.13, "pp0", {"locale": "zh"}), "13个百分点")
        self.assertEqual(format_value(0.13, "pp0", {"locale": "ms"}), f"13{NB}mata peratusan")
        self.assertEqual(format_value(1.6, "x1", {"locale": "vi"}), "1,6 lần")
        self.assertEqual(format_value(1.6, "x1", {"locale": "zh"}), "1.6倍")

    def test_dates(self):
        self.assertEqual(fmt_date("2026-10-20", "ru"), "20.10.2026")
        self.assertEqual(fmt_date("2026-10-20", "en"), "October 20, 2026")
        self.assertEqual(fmt_date("2026-10-20", "zh"), "2026年10月20日")
        self.assertEqual(fmt_date("2026-10-20", "vi"), "ngày 20 tháng 10 năm 2026")
        self.assertEqual(fmt_date("2026-10-20", "id"), "20 Oktober 2026")
        self.assertEqual(fmt_date("2026-10-20", "ms"), "20 Oktober 2026")

    def test_labels_complete(self):
        keys = set(LOCALES["en"]["labels"])
        for loc in LOCS:
            missing = keys - set(LOCALES[loc]["labels"])
            self.assertFalse(missing, f"{loc}: {missing}")

    def test_market_mapping(self):
        self.assertEqual(norm_market("vi"), "vn")
        self.assertEqual(norm_market("ms"), "my")
        self.assertEqual(norm_market("id"), "id")
        self.assertEqual(norm_market("zh"), "zh")

    def test_placeholders(self):
        self.assertEqual(resolve_placeholders("x {{a.b|pct0}}", {"a": {"b": 0.5}}), "x 50%")
        self.assertEqual(format_value(2030, "auto"), "2030")


class TestCitations(unittest.TestCase):
    SRC = {"id": "feld", "type": "web", "authors": ["Feld B."], "title": "The Rule of 40% for a healthy SaaS company",
           "container": "Feld Thoughts", "date": "2015-02-03", "url": "https://feld.com/archives/2015/02/rule-40-healthy-saas-company/",
           "accessed": "2026-10-20"}

    def test_gost(self):
        r = format_reference(self.SRC, "gost")
        self.assertIn("[Электронный ресурс]", r)
        self.assertIn("(дата обращения: 20.10.2026)", r)

    def test_apa(self):
        r = format_reference(self.SRC, "apa")
        self.assertTrue(r.startswith("Feld, B. (2015, February 3)."), r)
        self.assertNotIn("..", r)
        self.assertEqual(inline_citation([self.SRC], "apa"), "(Feld, 2015)")

    def test_chicago(self):
        r = format_reference(self.SRC, "chicago")
        self.assertIn("Accessed October 20, 2026", r)

    def test_gbt7714(self):
        r = format_reference(self.SRC, "gbt7714")
        self.assertIn("[EB/OL]", r)
        self.assertIn("(2015-02-03)[2026-10-20]", r)
        self.assertNotIn("..", r)
        law = {"id": "pipl", "type": "law", "title": "中华人民共和国个人信息保护法", "url": "http://www.npc.gov.cn/", "accessed": "2026-10-20"}
        self.assertIn("[EB/OL]", format_reference(law, "gbt7714"))

    def test_default_style_per_locale(self):
        self.assertEqual(LOCALES["ru"]["citation_style"], "footnotes")
        self.assertEqual(LOCALES["zh"]["citation_style"], "numeric")
        for loc in ("en", "vi", "id", "ms"):
            self.assertEqual(LOCALES[loc]["citation_style"], "author_date", loc)


class TestTextfitCJK(unittest.TestCase):
    def test_units(self):
        self.assertEqual(text_units("市场规模"), 2)   # 4 CJK chars ≈ 2 word units
        self.assertEqual(text_units("market size"), 2)

    def test_cjk_wraps_and_kinsoku(self):
        s = "这是一个用于测试中文换行规则的句子，标点不能出现在行首。"
        self.assertGreater(lines_needed(s, 18, 2.0), 1)
        self.assertEqual(lines_needed("短句", 18, 6.0), 1)
        from textfit import _tokens
        toks = _tokens("规则，标点")
        self.assertFalse(any(t.startswith("，") for t in toks), toks)

    def test_vietnamese_is_not_wide(self):
        self.assertEqual(lines_needed("Thị trường", 18, 6.0), 1)


# ---------------------------------------------------------------- market, finance, scoring
class TestMarketFinance(unittest.TestCase):
    def test_top_down_bottom_up(self):
        td = top_down(1000, [{"name": "a", "share": 0.5, "level": "SAM"}, {"name": "b", "share": 0.1, "level": "SOM"}])
        bu = bottom_up([{"name": "s", "units": 10, "arpa": 5, "sam_share": 0.5, "som_share": 0.1}])
        self.assertAlmostEqual(bu["TAM"], 50)
        self.assertAlmostEqual(bu["SAM"], 25)
        self.assertAlmostEqual(bu["SOM"], 2.5)
        self.assertTrue(triangulate(td, bu))

    def test_examples_all_locales(self):
        for loc in LOCS:
            r = size_market(load_json(os.path.join(EX(loc), "market_input.json")))
            self.assertIn("SAM", json.dumps(r), loc)
            f = fin_build(load_json(os.path.join(EX(loc), "finance_input.json")))
            self.assertIn("scenarios", f, loc)

    def test_same_economics_across_currencies(self):
        """Examples are the same fictional business scaled by an FX factor: MOIC must be identical."""
        moics = {loc: round(fin_build(load_json(os.path.join(EX(loc), "finance_input.json")))["base"]["investor"]["moic"], 6)
                 for loc in LOCS}
        self.assertEqual(len(set(moics.values())), 1, moics)


class TestScoring(unittest.TestCase):
    def setUp(self):
        self.m = load_model()
        self.stages = [k for k in self.m["stage_profiles"] if not k.startswith("_")]
        self.auds = [k for k in self.m["audience_multipliers"] if not k.startswith("_")] + ["neutral"]

    def test_weights_sum_100(self):
        self.assertEqual(sum(c["weight"] for c in self.m["criteria"]), 100)
        for st in self.stages:
            self.assertAlmostEqual(sum(self.m["stage_profiles"][st].values()), 100, places=6, msg=st)
        for au, w in self.m["deck_weights"].items():
            if not au.startswith("_"):
                self.assertAlmostEqual(sum(w.values()), 100, places=6, msg=au)

    def test_all_combinations_with_markets_sum_100(self):
        for mk in [None] + sorted(MARKETS):
            for st in self.stages:
                for au in self.auds:
                    w = weights_for(self.m, st, au, market=mk)
                    self.assertAlmostEqual(sum(w.values()), 100, places=4, msg=f"{mk}/{st}/{au}")

    def test_anchors_and_names(self):
        for c in self.m["criteria"]:
            for lang in ("ru", "en", "zh"):
                self.assertEqual(len(c[f"anchors_{lang}"]), 6, c["id"])
            for lang in LOCS:
                self.assertTrue(c.get(f"name_{lang}"), f"{c['id']} name_{lang}")

    def test_market_profiles(self):
        mps = self.m["market_profiles"]
        for mk in ("ru", "en", "zh", "vn", "id", "my"):
            p = mps[mk]
            reg = p["criteria"]["regulatory"]
            self.assertTrue(reg.get("definition_en"), mk)
            self.assertEqual(len(reg["anchors_en"]), 6, mk)
            self.assertTrue(os.path.exists(os.path.join(ROOT, p["guide"])), p["guide"])

    def test_example_scores_localized(self):
        for loc in LOCS:
            r = score(load_json(os.path.join(EX(loc), "scoring_input.json")))
            self.assertTrue(0 <= r["total"] <= 100)
            lo, hi = r["range"]
            self.assertLessEqual(lo, r["total"])
            self.assertGreaterEqual(hi, r["total"])
            self.assertEqual(r["market"], norm_market(loc))
            self.assertTrue(r["regulatory_definition"], loc)
            self.assertFalse([f for f in r["red_flags"] if f["severity"] == "high"], (loc, r["red_flags"]))

    def test_market_specific_red_flags(self):
        base = load_json(os.path.join(EX("id"), "scoring_input.json"))
        p = dict(base, audience="government", facts=dict(base["facts"], tkdn_certified=False, pse_registered=False))
        ids = {f["id"] for f in score(p)["red_flags"]}
        self.assertTrue({"id_tkdn", "id_pse"} <= ids, ids)
        p_vn = dict(load_json(os.path.join(EX("vi"), "scoring_input.json")))
        p_vn["facts"] = dict(p_vn["facts"], pdp_impact_assessment=False)
        self.assertIn("vn_pdp_dossier", {f["id"] for f in score(p_vn)["red_flags"]})
        p_my = dict(load_json(os.path.join(EX("ms"), "scoring_input.json")))
        p_my["facts"] = dict(p_my["facts"], pdpa_compliant=False)
        flags = score(p_my)["red_flags"]
        self.assertIn("my_pdpa", {f["id"] for f in flags})
        self.assertTrue(any("PDPA" in f["message"] for f in flags))
        # rules are filtered by market: a Russian-registry flag never fires for Malaysia
        p_my["facts"]["registry_eligible"] = False
        p_my["audience"] = "government"
        self.assertNotIn("registry_gov", {f["id"] for f in score(p_my)["red_flags"]})

    def test_perfect_score(self):
        payload = load_json(os.path.join(EX("en"), "scoring_input.json"))
        payload["scores"] = {c["id"]: {"score": 5, "evidence": "A"} for c in self.m["criteria"]}
        self.assertAlmostEqual(score(payload)["total"], 100, places=4)

    def test_allocate(self):
        self.assertTrue(allocate_slides(self.m, "investor", 16))

    def test_cli(self):
        out = run_cli(os.path.join(TOOLS, "scoring.py"), "profiles")
        self.assertTrue(json.loads(out.stdout))
        out = run_cli(os.path.join(TOOLS, "scoring.py"), "--input", os.path.join(EX("ms"), "scoring_input.json"),
                      "--market", "my")
        self.assertEqual(json.loads(out.stdout)["market"], "my")


# ---------------------------------------------------------------- full pipeline per locale
class TestPipelineAllLocales(unittest.TestCase):
    """Copies each example to a temp dir, runs the pipeline without rendering and checks the outputs."""

    @classmethod
    def setUpClass(cls):
        import run_pipeline
        cls.tmp = tempfile.mkdtemp()
        cls.out = {}
        for loc in LOCS:
            d = os.path.join(cls.tmp, loc)
            os.makedirs(d)
            for f in os.listdir(EX(loc)):
                if f.endswith(".json"):
                    shutil.copy(os.path.join(EX(loc), f), d)
            pj = os.path.join(d, "project.json")
            p = load_json(pj)
            p["render"] = False
            with open(pj, "w", encoding="utf8") as fh:
                json.dump(p, fh, ensure_ascii=False)
            run_pipeline.run(pj)
            cls.out[loc] = os.path.join(d, "output")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_qa_zero_errors(self):
        for loc in LOCS:
            s = load_json(os.path.join(self.out[loc], "pipeline_summary.json"))
            for theme in ("light", "dark"):
                self.assertTrue(s["qa"][theme]["pass"], (loc, theme))
                self.assertEqual(s["qa"][theme]["errors"], 0, (loc, theme))

    def test_native_charts_and_no_placeholders(self):
        for loc in LOCS:
            pptx = os.path.join(self.out[loc], "deck.pptx")
            docx = os.path.join(self.out[loc], "report.docx")
            with zipfile.ZipFile(pptx) as z:
                names = z.namelist()
            charts = [n for n in names if n.startswith("ppt/charts/chart") and n.endswith(".xml")]
            self.assertGreaterEqual(len(charts), 3, loc)
            self.assertEqual(len(charts), len([n for n in names if n.startswith("ppt/embeddings/")]), loc)
            self.assertFalse([n for n in names if n.startswith("ppt/media/")], loc)
            self.assertNotIn("{{", xml_text(pptx, "ppt/slides/"), loc)
            self.assertNotIn("{{", xml_text(docx, "word/document"), loc)

    def test_localized_numbers_in_outputs(self):
        expect = {"ru": ["млрд", "₽"], "en": ["$", "B"], "zh": ["亿元", "万元"], "vi": ["tỷ", "₫"],
                  "id": ["Rp", "miliar"], "ms": ["RM", "juta"]}
        for loc, needles in expect.items():
            txt = xml_text(os.path.join(self.out[loc], "deck.pptx"), "ppt/slides/")
            doc = xml_text(os.path.join(self.out[loc], "report.docx"), "word/document")
            for n in needles:
                self.assertIn(n, txt, (loc, n))
                self.assertIn(n, doc, (loc, n))

    def test_localized_titles_and_captions(self):
        cap = {"ru": "Таблица 1", "en": "Table 1", "zh": "表1", "vi": "Bảng 1", "id": "Tabel 1", "ms": "Jadual 1"}
        for loc, c in cap.items():
            doc = re.sub(r"<[^>]+>", "", xml_text(os.path.join(self.out[loc], "report.docx"), "word/document"))
            self.assertIn(c, doc, loc)

    def test_citation_style_per_locale(self):
        doc = lambda loc: re.sub(r"<[^>]+>", "", xml_text(os.path.join(self.out[loc], "report.docx"), "word/"))
        self.assertIn("дата обращения", doc("ru"))
        self.assertIn("[EB/OL]", doc("zh"))
        self.assertIn("(Feld, 2015)", doc("en"))
        for loc in ("vi", "id", "ms"):
            self.assertIn("(Feld, 2015)", doc(loc), loc)

    def test_cjk_fonts_in_xml(self):
        ppt = xml_text(os.path.join(self.out["zh"], "deck.pptx"), "ppt/slides/")
        self.assertRegex(ppt, r'<a:ea typeface="(Microsoft YaHei|Noto Sans CJK SC)"')
        doc = xml_text(os.path.join(self.out["zh"], "report.docx"), "word/")
        self.assertRegex(doc, r'w:eastAsia="(SimSun|SimHei|Noto Serif CJK SC|Noto Sans CJK SC)"')
        self.assertIn('w:eastAsia="zh-CN"', doc)
        charts = xml_text(os.path.join(self.out["zh"], "deck.pptx"), "ppt/charts/")
        self.assertIn("<a:ea ", charts)

    def test_theme_boilerplate_cleaned(self):
        """Only the Simplified Chinese script font remains in themes; East Asian theme language matches the locale."""
        for loc in LOCS:
            for f in ("deck.pptx", "report.docx"):
                xml = xml_text(os.path.join(self.out[loc], f), "")
                self.assertNotRegex(xml, r'<a:font script="(?!Hans")', (loc, f))
            settings = xml_text(os.path.join(self.out[loc], "report.docx"), "word/settings")
            m = re.search(r'<w:themeFontLang\b[^>]*w:eastAsia="([^"]+)"', settings)
            self.assertEqual(m.group(1) if m else None, "zh-CN" if loc == "zh" else None, loc)

    def test_core_properties_neutral(self):
        for loc in LOCS:
            for f in ("deck.pptx", "report.docx"):
                core = xml_text(os.path.join(self.out[loc], f), "docProps/core")
                self.assertIn("Velordoc", core, (loc, f))
                self.assertNotRegex(core, r"/(workspace|home)/", (loc, f))
                # no library-template defaults (template author, "generated by/using ...")
                self.assertNotRegex(core, r"generated (by|using) python", (loc, f))
                creator = re.search(r"<dc:creator>([^<]*)<", core)
                modifier = re.search(r"<cp:lastModifiedBy>([^<]*)<", core)
                if modifier:
                    self.assertEqual(modifier.group(1), creator.group(1) if creator else "", (loc, f))

    def test_banned_phrases_detected_per_locale(self):
        samples = {"vi": "Giải pháp mang tính cách mạng, không rủi ro", "id": "Solusi revolusioner tanpa risiko",
                   "ms": "Penyelesaian revolusioner tanpa risiko", "zh": "颠覆性的零风险方案",
                   "ru": "Революционный продукт", "en": "A revolutionary product"}
        d = tempfile.mkdtemp()
        for loc, title in samples.items():
            deck = {"meta": {"locale": loc}, "slides": [{"layout": "bullets", "title": title + " 2026",
                                                         "bullets": ["a", "b"], "notes": "n"}]}
            p = os.path.join(d, f"{loc}.pptx")
            build_pptx(deck, p, "light")
            r = run_qa(p, None, locale=loc)
            self.assertTrue(any(f["check"] == "banned_phrase" for f in r["findings"]), (loc, r["findings"]))

    def test_qa_detects_bad_deck(self):
        bad = {"meta": {"locale": "en"}, "slides": [
            {"layout": "chart", "title": "Market",
             "chart": {"type": "column", "categories": ["2025", "2026"], "series": [{"name": "x", "values": [1, 2]}]}}]}
        p = os.path.join(tempfile.mkdtemp(), "bad.pptx")
        build_pptx(bad, p, "light")
        self.assertFalse(run_qa(p, None, locale="en")["summary"]["pass"])

    def test_missing_cjk_font_is_flagged(self):
        deck = {"meta": {"locale": "en"}, "slides": [{"layout": "bullets", "title": "中文标题测试，市场规模达到三十亿元",
                                                      "bullets": ["中文内容"], "notes": "n"}]}
        p = os.path.join(tempfile.mkdtemp(), "cjk.pptx")
        build_pptx(deck, p, "light")
        from pptx import Presentation
        prs = Presentation(p)
        for sh in prs.slides[0].shapes:   # strip East Asian font slots to simulate a broken deck
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        for ea in run._r.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ea"):
                            ea.getparent().remove(ea)
                        run.font.name = "Arial"
        prs.save(p)
        r = run_qa(p, None, locale="zh")
        self.assertTrue(any(f["check"] == "cjk_font" for f in r["findings"]), r["findings"])


class TestScaffold(unittest.TestCase):
    def test_new_project_every_locale(self):
        import new_project
        for loc in LOCS:
            d = tempfile.mkdtemp()
            r = new_project.scaffold(d, "Test product", loc)
            self.assertEqual(r["locale"], loc)
            p = load_json(os.path.join(d, "project.json"))
            self.assertEqual(p["locale"], loc)
            self.assertEqual(p["out_dir"], "output")
            shutil.rmtree(d, ignore_errors=True)

    def test_scaffold_pipeline_runs(self):
        import new_project
        import run_pipeline
        d = tempfile.mkdtemp()
        new_project.scaffold(d, "Test", "ms")
        pj = os.path.join(d, "project.json")
        p = load_json(pj)
        p["render"] = False
        with open(pj, "w", encoding="utf8") as fh:
            json.dump(p, fh, ensure_ascii=False)
        run_pipeline.run(pj)
        s = load_json(os.path.join(d, "output", "pipeline_summary.json"))
        self.assertTrue(s["qa"]["light"]["pass"])
        shutil.rmtree(d, ignore_errors=True)

    def test_markdown_blocks(self):
        from build_docx import markdown_to_blocks
        b = markdown_to_blocks("# Summary\n\nRevenue grows.\n\n- a\n- b\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
        self.assertTrue(any(x["type"] == "heading" for x in b))
        self.assertTrue(any(x["type"] == "table" for x in b))


class TestRepositoryHygiene(unittest.TestCase):
    def test_required_files(self):
        for f in ("LICENSE", "README.md", "README.ru.md", "README.zh-CN.md", "SKILL.md", "CONTRIBUTING.md",
                  "CHANGELOG.md", "pyproject.toml", "requirements.txt", ".gitignore",
                  ".github/workflows/tests.yml", "research/presentation_techniques.md",
                  "scoring_model.md", "scoring_model.ru.md", "scoring_model.zh-CN.md"):
            self.assertTrue(os.path.exists(os.path.join(ROOT, f)), f)
        for mk in ("ru", "en", "zh", "vn", "id", "my"):
            self.assertTrue(os.path.exists(os.path.join(ROOT, "markets", f"{mk}.md")), mk)
        for loc in LOCS:
            self.assertEqual(len(os.listdir(os.path.join(ROOT, "prompts", loc))), 6, loc)
            self.assertEqual(len(os.listdir(os.path.join(ROOT, "templates", loc))), 6, loc)

    def test_only_supported_markets(self):
        """Exactly six market profiles and guides; no leftover profiles for unsupported markets."""
        model = json.load(open(os.path.join(ROOT, "config", "scoring_model.json"), encoding="utf-8"))
        self.assertEqual({k for k in model["market_profiles"] if not k.startswith("_")}, {"ru", "en", "zh", "vn", "id", "my"})
        self.assertEqual(sorted(os.listdir(os.path.join(ROOT, "markets"))),
                         sorted(["README.md", "ru.md", "en.md", "zh.md", "vn.md", "id.md", "my.md"]))
        self.assertEqual(sorted(os.listdir(os.path.join(ROOT, "prompts"))), sorted(LOCS))
        self.assertEqual(sorted(os.listdir(os.path.join(ROOT, "templates"))), sorted(LOCS))

    def test_banned_phrase_lists(self):
        b = load_json(os.path.join(ROOT, "config", "banned_phrases.json"))
        for loc in LOCS:
            self.assertGreaterEqual(len(b[loc]), 10, loc)
            self.assertIn(loc, b["allow_if_sourced"])

    def test_no_absolute_paths_in_text_files(self):
        pat = re.compile(r"/(workspace|home/[a-z]+)/")
        for dp, _, fs in os.walk(ROOT):
            if ".git" in dp or "__pycache__" in dp:
                continue
            for f in fs:
                if f.endswith((".md", ".json", ".py", ".toml", ".yml", ".txt")):
                    with open(os.path.join(dp, f), encoding="utf8", errors="ignore") as fh:
                        self.assertIsNone(pat.search(fh.read()), os.path.join(dp, f))


if __name__ == "__main__":
    unittest.main(verbosity=2)
